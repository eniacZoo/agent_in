#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""读取 .xlsx 第一张表的前若干行。优先 openpyxl（vendor/），否则 stdlib zipfile+xml。"""
import json
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
MAX_ROWS = 40
MAX_COLS = 16


def _parse_args():
    if "--args" in sys.argv:
        i = sys.argv.index("--args")
        if i + 1 < len(sys.argv):
            try:
                return json.loads(sys.argv[i + 1])
            except json.JSONDecodeError:
                return {}
    return {}


def _col_index(cell_ref):
    n = 0
    for ch in cell_ref:
        if not ch.isalpha():
            break
        n = n * 26 + (ord(ch.upper()) - ord("A") + 1)
    return max(n - 1, 0)


def _load_shared(z):
    names = [n for n in z.namelist() if n.lower().endswith("sharedstrings.xml")]
    if not names:
        return []
    root = ET.fromstring(z.read(names[0]))
    out = []
    for si in root.findall(NS + "si"):
        texts = [t.text or "" for t in si.iter(NS + "t")]
        out.append("".join(texts))
    return out


def _cell_text(c, shared):
    t = c.get("t")
    if t == "inlineStr":
        is_el = c.find(NS + "is")
        if is_el is None:
            return ""
        return "".join((x.text or "") for x in is_el.iter(NS + "t"))
    v = c.find(NS + "v")
    if v is None or v.text is None:
        return ""
    if t == "s":
        try:
            return shared[int(v.text)]
        except (ValueError, IndexError):
            return v.text
    return v.text


def _read_openpyxl(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows_out = []
        for i, row in enumerate(wb.active.iter_rows(values_only=True)):
            if i >= MAX_ROWS:
                break
            vals = [("" if c is None else str(c)) for c in list(row)[:MAX_COLS]]
            while vals and vals[-1] == "":
                vals.pop()
            if vals:
                rows_out.append("\t".join(vals))
        if not rows_out:
            return "(空表)"
        extra = "\n... (只显示前 {} 行)".format(MAX_ROWS) if len(rows_out) >= MAX_ROWS else ""
        return "\n".join(rows_out) + extra
    finally:
        wb.close()


def read_xlsx(path):
    p = Path(path)
    if not p.is_file():
        return "Error: 文件不存在: {}".format(path)
    if p.suffix.lower() != ".xlsx":
        return "Error: 只支持 .xlsx: {}".format(path)
    try:
        import openpyxl  # noqa: F401
        return _read_openpyxl(path)
    except ImportError:
        pass
    try:
        z = zipfile.ZipFile(p)
    except zipfile.BadZipFile:
        return "Error: 不是有效的 xlsx: {}".format(path)
    sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")]
    if not sheets:
        return "Error: 未找到工作表"
    sheets.sort()
    shared = _load_shared(z)
    root = ET.fromstring(z.read(sheets[0]))
    rows_out = []
    for row in root.findall(".//" + NS + "row"):
        cells = {}
        max_i = -1
        for c in row.findall(NS + "c"):
            ref = c.get("r") or ""
            i = _col_index(ref)
            cells[i] = _cell_text(c, shared)
            if i > max_i:
                max_i = i
        if max_i < 0:
            continue
        width = min(max_i + 1, MAX_COLS)
        line = [cells.get(i, "") for i in range(width)]
        rows_out.append("\t".join(line))
        if len(rows_out) >= MAX_ROWS:
            break
    z.close()
    if not rows_out:
        return "(空表)"
    extra = ""
    if len(rows_out) >= MAX_ROWS:
        extra = "\n... (只显示前 {} 行)".format(MAX_ROWS)
    return "\n".join(rows_out) + extra


def main():
    args = _parse_args()
    path = (args.get("path") or "").strip()
    if not path:
        sys.stdout.write("用法: /use read_xlsx <path.xlsx>\n")
        return
    sys.stdout.write(read_xlsx(path) + "\n")


if __name__ == "__main__":
    main()
