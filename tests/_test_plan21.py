#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""plan21.0：会话回写、运行时护栏、窗口压缩、glob/grep。"""
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, _ROOT)

import config
import context
import llm
import loop
import session
import tools


def _fake_chat_write_then_done(n):
    def fake_chat(messages, tools=None, stream=True, model=None, max_tokens=None, images=None):
        n["i"] += 1
        if n["i"] == 1:
            yield {
                "type": "tool_call",
                "id": "c1",
                "name": "write_file",
                "arguments": {"path": "temp/x.txt", "content": "hello"},
            }
            yield {"type": "usage", "data": {"prompt_tokens": 10, "completion_tokens": 2}}
        else:
            yield {"type": "text", "content": "wrote temp/x.txt"}
            yield {"type": "usage", "data": {"prompt_tokens": 12, "completion_tokens": 4}}
    return fake_chat


class Phase1SessionTests(unittest.TestCase):
    def test_transcript_keeps_tools(self):
        td = tempfile.mkdtemp(prefix="p21_tr_")
        tools.WORK_DIR = td
        tools.SAFE_MODE = True
        n = {"i": 0}
        old = llm.chat
        llm.chat = _fake_chat_write_then_done(n)
        buf = io.StringIO()
        try:
            with redirect_stdout(buf):
                text, *_rest, full = loop.agent_loop(
                    [{"role": "user", "content": "write x"}],
                    td, verbose=False, interactive=False,
                )
        finally:
            llm.chat = old
            shutil.rmtree(td, ignore_errors=True)
        roles = [m.get("role") for m in full]
        self.assertIn("tool", roles)
        self.assertTrue(any(m.get("role") == "assistant" and m.get("tool_calls") for m in full))
        self.assertIn("wrote temp/x.txt", text)
        self.assertEqual(sum(1 for m in full if m.get("content") == "wrote temp/x.txt"), 1)

    def test_validate_keeps_complete_tool_batch(self):
        msgs = [
            {"role": "user", "content": "go"},
            {
                "role": "assistant",
                "content": None,
                "tool_calls": [{
                    "id": "c1", "type": "function",
                    "function": {"name": "read_file", "arguments": "{}"},
                }],
            },
            {"role": "tool", "tool_call_id": "c1", "content": "ok"},
        ]
        out = session._validate_messages(msgs)
        self.assertEqual(len(out), 3)
        self.assertEqual(out[-1]["role"], "tool")

    def test_validate_truncates_orphan_tool_call(self):
        msgs = [
            {"role": "user", "content": "go"},
            {
                "role": "assistant",
                "content": None,
                "tool_calls": [{
                    "id": "c1", "type": "function",
                    "function": {"name": "read_file", "arguments": "{}"},
                }],
            },
        ]
        out = session._validate_messages(msgs)
        self.assertEqual(len(out), 1)
        self.assertEqual(out[0]["role"], "user")

    def test_no_duplicate_final_reply(self):
        src = (Path(_ROOT) / "agent.py").read_text(encoding="utf-8")
        self.assertNotIn('saved.append({"role": "assistant", "content": reply})', src)
        self.assertIn("session.absorb_path", src)


class Phase2HarnessTests(unittest.TestCase):
    def test_read_xlsx_bounded(self):
        td = tempfile.mkdtemp(prefix="p21_xlsx_")
        tools.WORK_DIR = td
        xlsx = os.path.join(td, "t.xlsx")
        tools._ensure_vendor()
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SheetA"
        ws["A1"] = "h1"
        ws["B1"] = "h2"
        wb.save(xlsx)
        wb.close()
        try:
            out = tools._exec_read_file({"path": xlsx})
            self.assertIn("SheetA", out)
            self.assertNotIn("16383", out)
            src = (Path(_ROOT) / "tools.py").read_text(encoding="utf-8")
            self.assertNotIn("range(1, ws.max_column", src)
        finally:
            shutil.rmtree(td, ignore_errors=True)

    def test_read_xlsx_not_raw_zip(self):
        td = tempfile.mkdtemp(prefix="p21_zip_")
        tools.WORK_DIR = td
        xlsx = os.path.join(td, "t.xlsx")
        tools._ensure_vendor()
        import openpyxl
        wb = openpyxl.Workbook()
        wb.save(xlsx)
        wb.close()
        try:
            out = tools._exec_read_file({"path": xlsx})
            self.assertFalse(out.lstrip().startswith("PK"))
            self.assertFalse(out.lstrip().startswith("xlsx") and out.startswith("PK"))
        finally:
            shutil.rmtree(td, ignore_errors=True)

    def test_temp_write_not_counted(self):
        wd = tempfile.mkdtemp(prefix="p21_wd_")
        try:
            self.assertFalse(loop._is_product_write("temp/b9.py", wd))
            self.assertFalse(loop._is_product_write(os.path.join(wd, "temp", "a.py"), wd))
            self.assertTrue(loop._is_product_write("out.xlsx", wd))
            self.assertTrue(loop._is_product_write(os.path.join(wd, "out.xlsx"), wd))
            self.assertTrue(loop._is_product_write(
                os.path.join(os.path.dirname(wd), "报告", "index.html"), wd))
            self.assertTrue(loop._is_temp_path("temp/explore.py", wd))
            self.assertFalse(loop._is_temp_path("index.html", wd))
        finally:
            shutil.rmtree(wd, ignore_errors=True)

    def test_blocked_write_not_counted(self):
        self.assertFalse(loop._is_tool_ok("操作被拒绝（blocked）：工作目录外，SAFE_MODE 开启"))
        self.assertFalse(loop._is_tool_ok("Error: boom"))
        self.assertTrue(loop._is_tool_ok("OK: Wrote x"))

    def test_stall_skips_inspect_only(self):
        self.assertFalse(loop.should_stall_stop(False, 12))
        self.assertTrue(loop.should_stall_stop(True, 12))
        self.assertFalse(loop.should_stall_stop(True, 11))

    def test_stall_after_temp_mutate(self):
        self.assertTrue(loop.should_stall_stop(True, 12))
        self.assertTrue(loop._shell_mutates_temp(r"python E:\x\temp\b9_probe.py"))
        self.assertFalse(loop._shell_mutates_temp("Get-ChildItem $HOME"))

    def test_reject_max_column_loop(self):
        td = tempfile.mkdtemp(prefix="p21_rej_")
        tools.WORK_DIR = td
        tools.SAFE_MODE = True
        script = os.path.join(td, "bad.py")
        sentinel = os.path.join(td, "sentinel.txt")
        Path(script).write_text(
            "from pathlib import Path\n"
            "Path(r'%s').write_text('ran')\n"
            "ws = type('W', (), {'max_column': 10})()\n"
            "for c in range(1, ws.max_column+1):\n"
            "    pass\n" % sentinel.replace("\\", "\\\\"),
            encoding="utf-8",
        )
        try:
            r = tools.execute(
                "shell",
                {"command": '"%s" "%s"' % (sys.executable, script)},
                confirm_fn=lambda _p: True,
            )
            self.assertTrue(r.lstrip().startswith("Error"), r[:300])
            self.assertIn("max_column", r)
            self.assertFalse(os.path.exists(sentinel))
        finally:
            shutil.rmtree(td, ignore_errors=True)

    def test_pythonunbuffered_in_env(self):
        self.assertEqual(tools._shell_env().get("PYTHONUNBUFFERED"), "1")


class Phase3CompactTests(unittest.TestCase):
    def test_window_trim_does_not_mutate_transcript(self):
        tr = [
            {"role": "user", "content": "a" * 8000},
            {"role": "assistant", "content": "ok"},
            {"role": "user", "content": "recent"},
        ]
        orig = json.dumps(tr, ensure_ascii=False)
        win = context.window_from_transcript(tr, "sys", keep_recent_tokens=20)
        context.apply(win, used_tokens=50000, limit=196000, budget=40000)
        self.assertEqual(json.dumps(tr, ensure_ascii=False), orig)

    def test_cut_on_boundary(self):
        asst = {
            "role": "assistant",
            "content": None,
            "tool_calls": [{
                "id": "c1", "type": "function",
                "function": {"name": "read_file", "arguments": "{}"},
            }],
        }
        tool = {"role": "tool", "tool_call_id": "c1", "content": "result"}
        tr = (
            [{"role": "user", "content": "old " + ("x" * 4000)}]
            + [asst, tool]
            + [{"role": "user", "content": "now"}]
        )
        win = context.window_from_transcript(tr, "sys", keep_recent_tokens=50)
        ns = [m for m in win if m.get("role") != "system"]
        self.assertTrue(ns)
        if ns[0].get("name") == "context_summary":
            ns = ns[1:]
        self.assertNotEqual(ns[0].get("role"), "tool")

    def test_keep_recent_tokens_config(self):
        config._config_cache = None
        self.assertEqual(config._DEFAULTS["keep_recent_tokens"], 20000)
        self.assertIn("keep_recent_tokens", config._WHITELIST)
        self.assertEqual(config.load()["keep_recent_tokens"], 20000)
        self.assertEqual(context.SUMMARY_MAX_CHARS, 1500)

    def test_summary_stays_out_of_session_file(self):
        tr = [{"role": "user", "content": "z" * 4000} for _ in range(4)]
        tr.append({"role": "user", "content": "recent"})
        win = context.window_from_transcript(tr, "sys", keep_recent_tokens=30)
        self.assertTrue(any(m.get("name") == "context_summary" for m in win))
        saved = loop._for_save(tr)
        self.assertFalse(any(m.get("name") == "context_summary" for m in saved))


class Phase4SearchTests(unittest.TestCase):
    def test_glob_and_grep(self):
        td = tempfile.mkdtemp(prefix="p21_g_")
        tools.WORK_DIR = td
        Path(td, "a.xlsx").write_bytes(b"not-a-real-xlsx")
        Path(td, "code.py").write_text("import openpyxl\n", encoding="utf-8")
        try:
            g = tools.execute("glob", {"pattern": "*.xlsx"})
            self.assertIn("a.xlsx", g)
            self.assertFalse(g.lstrip().startswith("Error"))
            out = tools.execute("glob", {"pattern": "*.xlsx", "path": os.path.abspath(os.sep)})
            self.assertTrue(out.lstrip().startswith("Error"))
            self.assertIn("outside", out.lower())
            hits = tools.execute("grep", {"pattern": "openpyxl"})
            self.assertIn("openpyxl", hits)
            self.assertIn("code.py", hits)
        finally:
            shutil.rmtree(td, ignore_errors=True)

    def test_tool_order(self):
        names = [t["function"]["name"] for t in tools.TOOLS]
        self.assertEqual(
            names,
            ["read_file", "write_file", "edit_file", "shell", "view_image", "glob", "grep"],
        )


if __name__ == "__main__":
    unittest.main()
